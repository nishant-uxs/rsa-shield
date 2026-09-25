"""
RSA-SHIELD — Simple college project PPT + PDF generator
Dark cybersecurity aesthetic, concise slides.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.colors import Color, HexColor, white
from reportlab.lib.units import inch, mm
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

OUT = Path(__file__).resolve().parent
PPTX_PATH = OUT / "RSA-SHIELD.pptx"
PDF_PATH = OUT / "RSA-SHIELD.pdf"

# Colors
BG = RGBColor(0x0B, 0x0E, 0x14)
BG2 = RGBColor(0x11, 0x15, 0x1D)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)
TEXT = RGBColor(0xE8, 0xEE, 0xF7)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
DIM = RGBColor(0x64, 0x74, 0x8B)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xF8, 0x71, 0x71)
GREEN = RGBColor(0x34, 0xD3, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# ReportLab hex
H_BG = HexColor("#0B0E14")
H_BG2 = HexColor("#11151D")
H_CYAN = HexColor("#22D3EE")
H_VIOLET = HexColor("#A78BFA")
H_TEXT = HexColor("#E8EEF7")
H_MUTED = HexColor("#94A3B8")
H_DIM = HexColor("#64748B")
H_AMBER = HexColor("#F59E0B")
H_RED = HexColor("#F87171")
H_GREEN = HexColor("#34D399")


def set_slide_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, left, top, width, height, fill_color, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    # soften corners
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=TEXT, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.alignment = align
    return box


def add_bullets(slide, left, top, width, height, items, size=16, color=MUTED):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(8)
    return box


def add_footer(slide, n, total=14):
    add_text(slide, Inches(0.5), Inches(7.05), Inches(3), Inches(0.3), "RSA-SHIELD", size=10, color=DIM)
    add_text(
        slide,
        Inches(11.5),
        Inches(7.05),
        Inches(1.3),
        Inches(0.3),
        f"{n:02d} / {total:02d}",
        size=10,
        color=DIM,
        align=PP_ALIGN.RIGHT,
    )


def eyebrow(slide, text, top=Inches(0.4)):
    add_text(slide, Inches(0.55), top, Inches(10), Inches(0.3), text.upper(), size=11, bold=True, color=CYAN)


def title(slide, text, top=Inches(0.7)):
    add_text(slide, Inches(0.55), top, Inches(12), Inches(0.7), text, size=28, bold=True, color=TEXT)


def make_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 14

    # ---- 1 COVER ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    add_rect(s, Inches(8.2), Inches(1.2), Inches(4.4), Inches(4.8), BG2, CYAN)
    add_text(s, Inches(0.7), Inches(1.5), Inches(7), Inches(0.35), "SECURITY AUDIT PLATFORM", size=12, bold=True, color=CYAN)
    add_text(s, Inches(0.7), Inches(2.0), Inches(7.5), Inches(1), "RSA-SHIELD", size=48, bold=True, color=CYAN)
    add_text(
        s,
        Inches(0.7),
        Inches(3.1),
        Inches(7),
        Inches(0.8),
        "RSA Key Security Auditor for\nLegacy Critical Infrastructure",
        size=18,
        color=MUTED,
    )
    add_text(
        s,
        Inches(0.7),
        Inches(4.3),
        Inches(7),
        Inches(1.8),
        "Team: [TEAM NAME]\nMembers: [MEMBER 1] · [MEMBER 2] · [MEMBER 3] · [MEMBER 4]\nMentor: [MENTOR NAME]\nInstitution: [COLLEGE / UNIVERSITY]",
        size=13,
        color=MUTED,
    )
    add_text(s, Inches(8.5), Inches(2.8), Inches(3.8), Inches(1.2), "n = p × q", size=28, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_text(s, Inches(8.5), Inches(3.8), Inches(3.8), Inches(0.8), "Mathematical RSA\nKey Auditing", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(s, 1, total)

    # ---- 2 PROBLEM ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "The Problem")
    title(s, "RSA Keys Can Fail Before Encryption Does")
    flow = [
        ("1. Legacy Assets", "SCADA / substations retain old crypto"),
        ("2. RSA Keys", "Security depends on key material quality"),
        ("3. Weak Modulus", "Poor entropy → factorable n"),
        ("4. Exposure", "Factorization-based risk"),
    ]
    for i, (h, d) in enumerate(flow):
        x = Inches(0.55 + i * 3.15)
        add_rect(s, x, Inches(2.2), Inches(2.95), Inches(2.6), BG2, CYAN if i < 3 else RED)
        add_text(s, x + Inches(0.15), Inches(2.45), Inches(2.65), Inches(0.5), h, size=14, bold=True, color=CYAN if i < 3 else RED)
        add_text(s, x + Inches(0.15), Inches(3.1), Inches(2.65), Inches(1.4), d, size=13, color=MUTED)
    add_text(s, Inches(0.55), Inches(5.2), Inches(12), Inches(0.8), "Manual auditing of large key inventories cannot prove mathematical strength.", size=15, color=MUTED)
    add_footer(s, 2, total)

    # ---- 3 GAP ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Existing Gap")
    title(s, "From Key Inventory to Mathematical Assurance")
    add_rect(s, Inches(0.55), Inches(1.8), Inches(5.7), Inches(4.4), BG2, RED)
    add_text(s, Inches(0.8), Inches(2.0), Inches(5), Inches(0.4), "Traditional Inventory", size=16, bold=True, color=RED)
    add_bullets(
        s,
        Inches(0.8),
        Inches(2.6),
        Inches(5),
        Inches(3.2),
        ["Key inventory", "Basic length check", "Basic config check", "❌ No mathematical strength proof"],
        size=15,
        color=MUTED,
    )
    add_rect(s, Inches(7.0), Inches(1.8), Inches(5.7), Inches(4.4), BG2, CYAN)
    add_text(s, Inches(7.25), Inches(2.0), Inches(5), Inches(0.4), "RSA-SHIELD", size=16, bold=True, color=CYAN)
    add_bullets(
        s,
        Inches(7.25),
        Inches(2.6),
        Inches(5),
        Inches(3.2),
        ["Structural analysis", "Mathematical analysis", "ECM factorization attempt", "Evidence-based assessment"],
        size=15,
        color=MUTED,
    )
    add_footer(s, 3, total)

    # ---- 4 SOLUTION ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Our Solution")
    title(s, "RSA-SHIELD: Mathematical RSA Security Auditor")
    add_text(
        s,
        Inches(0.55),
        Inches(1.5),
        Inches(12),
        Inches(0.5),
        "Software platform that turns established crypto math into a practical audit tool.",
        size=15,
        color=MUTED,
    )
    steps = ["Key Input", "Parse", "Extract n,e", "Analyze", "ECM", "Evidence", "Risk", "Report"]
    for i, st in enumerate(steps):
        x = Inches(0.4 + i * 1.6)
        add_rect(s, x, Inches(2.6), Inches(1.45), Inches(1.5), BG2, CYAN)
        add_text(s, x, Inches(3.05), Inches(1.45), Inches(0.7), st, size=12, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            add_text(s, x + Inches(1.35), Inches(3.1), Inches(0.3), Inches(0.4), "→", size=14, color=DIM)
    add_text(
        s,
        Inches(0.55),
        Inches(4.8),
        Inches(12),
        Inches(1),
        "We did NOT invent RSA or ECM. We operationalize them for authorized key auditing.",
        size=14,
        color=VIOLET,
    )
    add_footer(s, 4, total)

    # ---- 5 HOW IT WORKS ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Audit Engine")
    title(s, "Inside the Audit Engine")
    add_rect(s, Inches(0.55), Inches(1.7), Inches(5.5), Inches(1.2), BG2, CYAN)
    add_text(s, Inches(0.8), Inches(1.95), Inches(5), Inches(0.7), "n = p × q\nRecovering p or q compromises the modulus.", size=15, color=TEXT)
    left = [
        "1. Accept public key / modulus",
        "2. Validate format",
        "3. Extract n and e",
        "4. Preliminary checks",
    ]
    right = [
        "5. Run configurable ECM",
        "6. Detect factors",
        "7. Estimate exposure",
        "8. Generate audit result",
    ]
    add_bullets(s, Inches(0.55), Inches(3.3), Inches(5.8), Inches(3), left, size=15)
    add_bullets(s, Inches(6.8), Inches(3.3), Inches(5.8), Inches(3), right, size=15)
    add_footer(s, 5, total)

    # ---- 6 ECM ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Cryptographic Foundation")
    title(s, "Why Elliptic Curve Method (ECM)?")
    add_bullets(
        s,
        Inches(0.55),
        Inches(1.7),
        Inches(6),
        Inches(2.5),
        [
            "ECM = established factorization algorithm",
            "Good at finding relatively small factors",
            "RSA: n = p × q",
            "A non-trivial factor = concrete weakness evidence",
        ],
        size=15,
    )
    add_rect(s, Inches(7.0), Inches(1.7), Inches(5.7), Inches(4.5), BG2, VIOLET)
    add_text(s, Inches(7.3), Inches(2.0), Inches(5.2), Inches(0.4), "ECM Audit Flow", size=14, bold=True, color=VIOLET)
    ecm_flow = "RSA modulus n\n↓\nECM search\n↓\nElliptic curve ops\n↓\nFactor found?  → YES: evidence\n               → NO: inconclusive"
    add_text(s, Inches(7.3), Inches(2.6), Inches(5.2), Inches(3.2), ecm_flow, size=14, color=MUTED)
    add_text(
        s,
        Inches(0.55),
        Inches(5.5),
        Inches(6),
        Inches(0.8),
        "ECM does NOT instantly factor modern strong RSA keys.",
        size=13,
        color=AMBER,
    )
    add_footer(s, 6, total)

    # ---- 7 ARCHITECTURE ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "System Design")
    title(s, "RSA-SHIELD Architecture")
    layers = [
        ("Frontend", "React + TypeScript + Tailwind"),
        ("API", "FastAPI + Python"),
        ("Engines", "RSA Parser · Key Analyzer · ECM · Risk · Reports"),
        ("Data", "PostgreSQL / SQLite"),
        ("Output", "JSON + PDF audit reports"),
    ]
    for i, (h, d) in enumerate(layers):
        y = Inches(1.65 + i * 0.85)
        add_rect(s, Inches(1.5), y, Inches(10.3), Inches(0.75), BG2, CYAN)
        add_text(s, Inches(1.8), y + Inches(0.18), Inches(2.5), Inches(0.45), h, size=14, bold=True, color=CYAN)
        add_text(s, Inches(4.5), y + Inches(0.18), Inches(7), Inches(0.45), d, size=14, color=MUTED)
    add_footer(s, 7, total)

    # ---- 8 LAYERS ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Analysis Engine")
    title(s, "More Than a Factorization Tool")
    layers = [
        "L1  Format & Parameter Validation",
        "L2  RSA Parameter Analysis",
        "L3  Modulus Characteristics",
        "L4  ECM Factorization Attempt",
        "L5  Evidence-Based Risk Classification",
        "L6  Audit Report",
    ]
    for i, t in enumerate(layers):
        y = Inches(1.6 + i * 0.7)
        add_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.6), BG2, CYAN if i % 2 == 0 else VIOLET)
        add_text(s, Inches(0.85), y + Inches(0.12), Inches(11.5), Inches(0.4), t, size=15, color=TEXT)
    add_text(
        s,
        Inches(0.55),
        Inches(6.0),
        Inches(12),
        Inches(0.5),
        "Separates mathematical evidence from heuristic warnings.",
        size=14,
        color=VIOLET,
    )
    add_footer(s, 8, total)

    # ---- 9 SCENARIO ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Audit Scenario")
    title(s, "From Suspicious Key to Security Evidence")
    path = "Legacy Device → Public Key → RSA-SHIELD → Suspicious Modulus → ECM → Factor Found → Alert → Key Rotation"
    add_text(s, Inches(0.55), Inches(1.55), Inches(12.2), Inches(0.6), path, size=13, color=MUTED)
    add_rect(s, Inches(0.55), Inches(2.3), Inches(12.2), Inches(4.0), BG2, RED)
    add_text(s, Inches(0.9), Inches(2.55), Inches(6), Inches(0.4), "Mock Audit Dashboard", size=14, bold=True, color=RED)
    rows = [
        ("Key ID", "LEGACY-SCADA-0042"),
        ("Modulus size", "1024-bit"),
        ("Public exponent", "e = 65537"),
        ("ECM status", "Factor discovered"),
        ("Severity", "CRITICAL"),
        ("Action", "Recommend key rotation"),
    ]
    for i, (k, v) in enumerate(rows):
        col = i % 2
        row = i // 2
        x = Inches(0.9 + col * 5.8)
        y = Inches(3.15 + row * 0.85)
        add_text(s, x, y, Inches(5), Inches(0.3), k, size=11, color=DIM)
        add_text(s, x, y + Inches(0.28), Inches(5), Inches(0.35), v, size=15, bold=True, color=TEXT if k != "Severity" else RED)
    add_footer(s, 9, total)

    # ---- 10 STACK ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Implementation")
    title(s, "Technology Stack")
    stack = [
        ("UI", "React · TypeScript · Tailwind · Vite · Recharts"),
        ("API", "Python · FastAPI · Pydantic"),
        ("Security Engine", "RSA math · Modular arithmetic · Lenstra ECM"),
        ("Data", "PostgreSQL / SQLite"),
        ("DevOps", "Docker · Git · GitHub"),
        ("Testing", "Pytest · Unit + Integration"),
    ]
    for i, (h, d) in enumerate(stack):
        y = Inches(1.6 + i * 0.75)
        add_rect(s, Inches(0.55), y, Inches(12.2), Inches(0.65), BG2, CYAN)
        add_text(s, Inches(0.85), y + Inches(0.15), Inches(2.8), Inches(0.4), h, size=14, bold=True, color=CYAN)
        add_text(s, Inches(3.8), y + Inches(0.15), Inches(8.5), Inches(0.4), d, size=14, color=MUTED)
    add_footer(s, 10, total)

    # ---- 11 FEATURES ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Product")
    title(s, "Key Features")
    feats = [
        "RSA public-key upload",
        "Batch key auditing",
        "Modulus analysis",
        "ECM factorization attempts",
        "Audit history",
        "Risk classification",
        "Evidence tracking",
        "Exportable reports (JSON/PDF)",
        "API access",
        "Audit simulation",
        "Crypto visualizations",
        "Clear evidence tiers",
    ]
    for i, f in enumerate(feats):
        col = i % 3
        row = i // 3
        x = Inches(0.55 + col * 4.2)
        y = Inches(1.7 + row * 1.15)
        add_rect(s, x, y, Inches(3.95), Inches(1.0), BG2, CYAN)
        add_text(s, x + Inches(0.2), y + Inches(0.3), Inches(3.55), Inches(0.5), f, size=14, color=TEXT)
    add_footer(s, 11, total)

    # ---- 12 EVALUATION ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Evaluation")
    title(s, "How We Measure the System")
    metrics = [
        "Keys successfully parsed",
        "Audit processing time",
        "ECM execution time",
        "Factor discovery on weak test keys",
        "False-positive rate (heuristics)",
        "Batch throughput",
        "Resource utilization",
        "Detection rate (lab weak keys only)",
    ]
    add_bullets(s, Inches(0.55), Inches(1.7), Inches(6.5), Inches(4.5), metrics, size=15)
    add_rect(s, Inches(7.3), Inches(2.2), Inches(5.3), Inches(2.8), BG2, AMBER)
    add_text(s, Inches(7.6), Inches(2.7), Inches(4.8), Inches(1.8), "Prototype Benchmark\n— To Be Measured —\n\nNo fabricated numbers.", size=16, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
    add_footer(s, 12, total)

    # ---- 13 LIMITATIONS ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Responsible Use")
    title(s, "What RSA-SHIELD Does — and Does Not — Claim")
    add_rect(s, Inches(0.55), Inches(1.7), Inches(5.9), Inches(4.5), BG2, CYAN)
    add_text(s, Inches(0.85), Inches(1.95), Inches(5.3), Inches(0.4), "Does", size=16, bold=True, color=CYAN)
    add_bullets(
        s,
        Inches(0.85),
        Inches(2.5),
        Inches(5.3),
        Inches(3.4),
        [
            "Audit authorized RSA keys",
            "Structural + math analysis",
            "ECM as evidence probe",
            "Evidence-tiered labels",
            "Human-readable reports",
        ],
        size=14,
    )
    add_rect(s, Inches(6.85), Inches(1.7), Inches(5.9), Inches(4.5), BG2, RED)
    add_text(s, Inches(7.15), Inches(1.95), Inches(5.3), Inches(0.4), "Does Not Claim", size=16, bold=True, color=RED)
    add_bullets(
        s,
        Inches(7.15),
        Inches(2.5),
        Inches(5.3),
        Inches(3.4),
        [
            "Not a magic break of modern RSA",
            "No factor ≠ proof of security",
            "Not a formal security proof",
            "Needs specialist validation",
            "Authorized keys only",
        ],
        size=14,
    )
    add_footer(s, 13, total)

    # ---- 14 FUTURE ----
    s = prs.slides.add_slide(blank)
    set_slide_bg(s)
    eyebrow(s, "Roadmap")
    title(s, "From RSA Auditor to Crypto Asset Intelligence")
    future = [
        "Large-scale key inventory",
        "Certificate scanning",
        "Weak-key pattern detection",
        "Historical monitoring",
        "PQC migration readiness",
        "Enterprise dashboards",
        "CI/CD crypto scanning",
        "Compliance reports",
        "Hardware key metadata",
    ]
    for i, f in enumerate(future):
        col = i % 3
        row = i // 3
        x = Inches(0.55 + col * 4.2)
        y = Inches(1.7 + row * 1.0)
        add_rect(s, x, y, Inches(3.95), Inches(0.85), BG2, VIOLET)
        add_text(s, x + Inches(0.2), y + Inches(0.25), Inches(3.55), Inches(0.45), f, size=13, color=TEXT)
    add_rect(s, Inches(0.55), Inches(5.0), Inches(12.2), Inches(1.5), BG2, CYAN)
    add_text(s, Inches(0.8), Inches(5.2), Inches(11.7), Inches(0.45), "Know Your Keys. Prove Their Strength.", size=20, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
    add_text(
        s,
        Inches(0.8),
        Inches(5.75),
        Inches(11.7),
        Inches(0.55),
        "[TEAM NAME] · [MEMBERS] · [COLLEGE] · GitHub: [REPO] · Demo: [URL]",
        size=12,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(s, 14, total)

    prs.save(PPTX_PATH)
    print(f"Saved: {PPTX_PATH}")


def make_pdf():
    page = landscape(A4)
    w, h = page
    c = canvas.Canvas(str(PDF_PATH), pagesize=page)
    total = 14

    def bg():
        c.setFillColor(H_BG)
        c.rect(0, 0, w, h, fill=1, stroke=0)

    def footer(n):
        c.setFillColor(H_DIM)
        c.setFont("Helvetica", 9)
        c.drawString(28, 18, "RSA-SHIELD")
        c.drawRightString(w - 28, 18, f"{n:02d} / {total:02d}")

    def card(x, y, cw, ch, stroke=H_CYAN):
        c.setFillColor(H_BG2)
        c.setStrokeColor(stroke)
        c.setLineWidth(1)
        c.roundRect(x, y, cw, ch, 8, fill=1, stroke=1)

    def ey(txt):
        c.setFillColor(H_CYAN)
        c.setFont("Helvetica-Bold", 10)
        c.drawString(36, h - 40, txt.upper())

    def ttl(txt):
        c.setFillColor(H_TEXT)
        c.setFont("Helvetica-Bold", 22)
        c.drawString(36, h - 72, txt)

    # 1 Cover
    bg()
    ey("Security Audit Platform")
    c.setFillColor(H_CYAN)
    c.setFont("Helvetica-Bold", 42)
    c.drawString(36, h - 130, "RSA-SHIELD")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 14)
    c.drawString(36, h - 160, "RSA Key Security Auditor for Legacy Critical Infrastructure")
    c.setFont("Helvetica", 11)
    for i, line in enumerate(
        [
            "Team: [TEAM NAME]",
            "Members: [MEMBER 1] · [MEMBER 2] · [MEMBER 3] · [MEMBER 4]",
            "Mentor: [MENTOR NAME]",
            "Institution: [COLLEGE / UNIVERSITY]",
        ]
    ):
        c.drawString(36, h - 220 - i * 22, line)
    card(w - 280, 120, 240, 220)
    c.setFillColor(H_CYAN)
    c.setFont("Helvetica-Bold", 22)
    c.drawCentredString(w - 160, 230, "n = p × q")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 11)
    c.drawCentredString(w - 160, 190, "Mathematical RSA")
    c.drawCentredString(w - 160, 172, "Key Auditing")
    footer(1)
    c.showPage()

    # 2 Problem
    bg()
    ey("The Problem")
    ttl("RSA Keys Can Fail Before Encryption Does")
    boxes = [
        ("1. Legacy Assets", "SCADA / old crypto", H_CYAN),
        ("2. RSA Keys", "Quality of key material", H_CYAN),
        ("3. Weak Modulus", "Poor entropy risk", H_AMBER),
        ("4. Exposure", "Factorization risk", H_RED),
    ]
    for i, (t, d, col) in enumerate(boxes):
        x = 36 + i * 190
        card(x, h - 320, 175, 140, col)
        c.setFillColor(col)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(x + 12, h - 210, t)
        c.setFillColor(H_MUTED)
        c.setFont("Helvetica", 10)
        c.drawString(x + 12, h - 235, d)
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    c.drawString(36, 80, "Manual inventory checks cannot prove mathematical strength.")
    footer(2)
    c.showPage()

    # 3 Gap
    bg()
    ey("Existing Gap")
    ttl("From Key Inventory to Mathematical Assurance")
    card(36, 80, 340, 280, H_RED)
    c.setFillColor(H_RED)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(56, 330, "Traditional Inventory")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    for i, t in enumerate(["Key inventory", "Length check", "Config check", "No math proof"]):
        c.drawString(56, 290 - i * 28, "•  " + t)
    card(400, 80, 340, 280, H_CYAN)
    c.setFillColor(H_CYAN)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(420, 330, "RSA-SHIELD")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    for i, t in enumerate(["Structural analysis", "Mathematical analysis", "ECM attempt", "Evidence-based report"]):
        c.drawString(420, 290 - i * 28, "•  " + t)
    footer(3)
    c.showPage()

    # 4 Solution
    bg()
    ey("Our Solution")
    ttl("RSA-SHIELD: Mathematical RSA Security Auditor")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    c.drawString(36, h - 100, "Operationalizes established crypto math into a practical audit tool.")
    steps = ["Input", "Parse", "n,e", "Analyze", "ECM", "Evidence", "Risk", "Report"]
    for i, st in enumerate(steps):
        x = 30 + i * 95
        card(x, h - 250, 85, 70, H_CYAN)
        c.setFillColor(H_CYAN)
        c.setFont("Helvetica-Bold", 10)
        c.drawCentredString(x + 42, h - 215, st)
    c.setFillColor(H_VIOLET)
    c.setFont("Helvetica", 12)
    c.drawString(36, 100, "We did NOT invent RSA or ECM — we apply them for authorized auditing.")
    footer(4)
    c.showPage()

    # 5 Engine
    bg()
    ey("Audit Engine")
    ttl("Inside the Audit Engine")
    card(36, h - 160, 360, 60, H_CYAN)
    c.setFillColor(H_TEXT)
    c.setFont("Helvetica-Bold", 14)
    c.drawString(52, h - 125, "n = p × q   →   recovering factors compromises RSA")
    left = ["Accept public key / modulus", "Validate format", "Extract n and e", "Preliminary checks"]
    right = ["Configurable ECM", "Detect factors", "Estimate exposure", "Generate audit result"]
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    for i, t in enumerate(left):
        c.drawString(50, h - 220 - i * 28, f"{i+1}.  {t}")
    for i, t in enumerate(right):
        c.drawString(420, h - 220 - i * 28, f"{i+5}.  {t}")
    footer(5)
    c.showPage()

    # 6 ECM
    bg()
    ey("Cryptographic Foundation")
    ttl("Why Elliptic Curve Method (ECM)?")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    for i, t in enumerate(
        [
            "ECM = established factorization algorithm",
            "Useful for finding relatively small factors",
            "RSA modulus: n = p × q",
            "Non-trivial factor = concrete weakness evidence",
        ]
    ):
        c.drawString(40, h - 120 - i * 28, "•  " + t)
    card(420, 90, 300, 250, H_VIOLET)
    c.setFillColor(H_VIOLET)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(440, 310, "ECM Audit Flow")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 11)
    for i, t in enumerate(["RSA modulus n", "↓ ECM search", "↓ Curve operations", "Factor? YES → evidence", "Factor? NO → inconclusive"]):
        c.drawString(440, 270 - i * 28, t)
    c.setFillColor(H_AMBER)
    c.setFont("Helvetica", 11)
    c.drawString(40, 70, "ECM does NOT instantly factor modern strong RSA keys.")
    footer(6)
    c.showPage()

    # 7 Architecture
    bg()
    ey("System Design")
    ttl("RSA-SHIELD Architecture")
    rows = [
        ("Frontend", "React + TypeScript + Tailwind"),
        ("API", "FastAPI + Python"),
        ("Engines", "Parser · Analyzer · ECM · Risk · Reports"),
        ("Data", "PostgreSQL / SQLite"),
        ("Output", "JSON + PDF audit reports"),
    ]
    for i, (a, b) in enumerate(rows):
        y = h - 140 - i * 55
        card(80, y - 15, 640, 45, H_CYAN)
        c.setFillColor(H_CYAN)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(100, y, a)
        c.setFillColor(H_MUTED)
        c.setFont("Helvetica", 12)
        c.drawString(260, y, b)
    footer(7)
    c.showPage()

    # 8 Layers
    bg()
    ey("Analysis Engine")
    ttl("More Than a Factorization Tool")
    layers = [
        "L1  Format & Parameter Validation",
        "L2  RSA Parameter Analysis",
        "L3  Modulus Characteristics",
        "L4  ECM Factorization Attempt",
        "L5  Evidence-Based Risk Classification",
        "L6  Audit Report",
    ]
    for i, t in enumerate(layers):
        y = h - 130 - i * 42
        card(50, y - 10, 700, 35, H_CYAN if i % 2 == 0 else H_VIOLET)
        c.setFillColor(H_TEXT)
        c.setFont("Helvetica", 12)
        c.drawString(70, y, t)
    c.setFillColor(H_VIOLET)
    c.setFont("Helvetica", 11)
    c.drawString(50, 50, "Separates mathematical evidence from heuristic warnings.")
    footer(8)
    c.showPage()

    # 9 Scenario
    bg()
    ey("Audit Scenario")
    ttl("From Suspicious Key to Security Evidence")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 10)
    c.drawString(36, h - 100, "Legacy Device → Key → RSA-SHIELD → Suspicious → ECM → Factor → Alert → Rotation")
    card(36, 70, 740, 250, H_RED)
    c.setFillColor(H_RED)
    c.setFont("Helvetica-Bold", 12)
    c.drawString(56, 290, "Mock Audit Dashboard")
    data = [
        ("Key ID", "LEGACY-SCADA-0042"),
        ("Modulus", "1024-bit"),
        ("Exponent", "e = 65537"),
        ("ECM status", "Factor discovered"),
        ("Severity", "CRITICAL"),
        ("Action", "Recommend key rotation"),
    ]
    for i, (k, v) in enumerate(data):
        col = i % 2
        row = i // 2
        x = 56 + col * 360
        y = 240 - row * 50
        c.setFillColor(H_DIM)
        c.setFont("Helvetica", 9)
        c.drawString(x, y, k)
        c.setFillColor(H_RED if k == "Severity" else H_TEXT)
        c.setFont("Helvetica-Bold", 12)
        c.drawString(x, y - 16, v)
    footer(9)
    c.showPage()

    # 10 Stack
    bg()
    ey("Implementation")
    ttl("Technology Stack")
    stack = [
        ("UI", "React · TypeScript · Tailwind · Vite"),
        ("API", "Python · FastAPI · Pydantic"),
        ("Engine", "RSA math · Modular arithmetic · Lenstra ECM"),
        ("Data", "PostgreSQL / SQLite"),
        ("DevOps", "Docker · Git · GitHub"),
        ("Testing", "Pytest · Unit + Integration"),
    ]
    for i, (a, b) in enumerate(stack):
        y = h - 130 - i * 45
        card(50, y - 10, 700, 38, H_CYAN)
        c.setFillColor(H_CYAN)
        c.setFont("Helvetica-Bold", 11)
        c.drawString(70, y, a)
        c.setFillColor(H_MUTED)
        c.setFont("Helvetica", 11)
        c.drawString(200, y, b)
    footer(10)
    c.showPage()

    # 11 Features
    bg()
    ey("Product")
    ttl("Key Features")
    feats = [
        "RSA key upload",
        "Batch auditing",
        "Modulus analysis",
        "ECM attempts",
        "Audit history",
        "Risk classification",
        "Evidence tracking",
        "JSON/PDF reports",
        "API access",
        "Audit simulation",
        "Visualizations",
        "Evidence tiers",
    ]
    for i, f in enumerate(feats):
        col = i % 3
        row = i // 3
        x = 40 + col * 250
        y = h - 160 - row * 70
        card(x, y - 15, 230, 50, H_CYAN)
        c.setFillColor(H_TEXT)
        c.setFont("Helvetica", 11)
        c.drawString(x + 15, y, f)
    footer(11)
    c.showPage()

    # 12 Evaluation
    bg()
    ey("Evaluation")
    ttl("How We Measure the System")
    metrics = [
        "Keys successfully parsed",
        "Audit processing time",
        "ECM execution time",
        "Factor discovery on weak test keys",
        "False-positive rate (heuristics)",
        "Batch throughput / resource use",
    ]
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 12)
    for i, t in enumerate(metrics):
        c.drawString(50, h - 120 - i * 28, "•  " + t)
    card(420, 120, 300, 180, H_AMBER)
    c.setFillColor(H_AMBER)
    c.setFont("Helvetica-Bold", 14)
    c.drawCentredString(570, 230, "Prototype Benchmark")
    c.drawCentredString(570, 205, "— To Be Measured —")
    c.setFont("Helvetica", 11)
    c.drawCentredString(570, 170, "No fabricated numbers.")
    footer(12)
    c.showPage()

    # 13 Limitations
    bg()
    ey("Responsible Use")
    ttl("What RSA-SHIELD Does — and Does Not — Claim")
    card(36, 70, 350, 280, H_CYAN)
    c.setFillColor(H_CYAN)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(56, 320, "Does")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 11)
    for i, t in enumerate(
        ["Audit authorized RSA keys", "Structural + math analysis", "ECM as evidence probe", "Evidence-tiered labels", "Readable reports"]
    ):
        c.drawString(56, 280 - i * 28, "•  " + t)
    card(410, 70, 350, 280, H_RED)
    c.setFillColor(H_RED)
    c.setFont("Helvetica-Bold", 13)
    c.drawString(430, 320, "Does Not Claim")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 11)
    for i, t in enumerate(
        ["Not a magic break of modern RSA", "No factor ≠ security proof", "Not a formal proof", "Needs specialist validation", "Authorized keys only"]
    ):
        c.drawString(430, 280 - i * 28, "•  " + t)
    footer(13)
    c.showPage()

    # 14 Future
    bg()
    ey("Roadmap")
    ttl("From RSA Auditor to Crypto Asset Intelligence")
    future = [
        "Large-scale inventory",
        "Certificate scanning",
        "Weak-key patterns",
        "Historical monitoring",
        "PQC readiness",
        "Enterprise dashboards",
        "CI/CD scanning",
        "Compliance reports",
        "Hardware metadata",
    ]
    for i, f in enumerate(future):
        col = i % 3
        row = i // 3
        x = 40 + col * 250
        y = h - 150 - row * 65
        card(x, y - 10, 230, 45, H_VIOLET)
        c.setFillColor(H_TEXT)
        c.setFont("Helvetica", 11)
        c.drawString(x + 15, y, f)
    card(36, 40, 740, 70, H_CYAN)
    c.setFillColor(H_CYAN)
    c.setFont("Helvetica-Bold", 16)
    c.drawCentredString(w / 2, 80, "Know Your Keys. Prove Their Strength.")
    c.setFillColor(H_MUTED)
    c.setFont("Helvetica", 10)
    c.drawCentredString(w / 2, 55, "[TEAM NAME] · [MEMBERS] · [COLLEGE] · GitHub / Demo placeholders")
    footer(14)
    c.showPage()

    c.save()
    print(f"Saved: {PDF_PATH}")


if __name__ == "__main__":
    make_pptx()
    make_pdf()
